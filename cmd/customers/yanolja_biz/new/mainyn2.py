import csv
import json
import time
import threading

from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


# ==================================================
# 설정
# ==================================================

INPUT_CSV = "yanolja_local_accommodation.csv"

OUTPUT_CSV = (
    "yanolja_local_accommodation_with_seller.csv"
)

OUTPUT_XLSX = (
    "yanolja_local_accommodation_with_seller.xlsx"
)


TRPC_URL = (
    "https://nol.yanolja.com/stay/api/trpc/"
    "stay.properties.getFavorite,"
    "stay.properties.getSellerInfo"
)


_print_lock = threading.Lock()


def safe_print(msg):
    with _print_lock:
        print(msg, flush=True)


# ==================================================
# 옵션
# ==================================================

def create_options():

    return {
        # 동시에 상세 몇 개 호출할지
        # 너무 에러 많이 나면 4~6으로 낮추기
        "maxWorkers": 6,

        # 숙소 처리 완료 로그 사이 간격
        "sleepSec": 0.05,

        # 실패 후 재시도 횟수
        # retry=3이면
        # 최초 + 재시도 3번까지 가능
        "retry": 3,

        # HTTP timeout
        "timeout": 20,

        # 로그 값 길이
        # 0이면 전체 출력
        "logValueLimit": 0,

        # 재시도 기본 대기시간
        "retrySleepSec": 1,
    }


# ==================================================
# CSV 읽기
# ==================================================

def read_csv_rows(path):

    rows = []

    with open(
            path,
            "r",
            encoding="utf-8-sig",
            newline=""
    ) as f:

        reader = csv.DictReader(f)

        for row in reader:
            rows.append(row)

    return rows


# ==================================================
# CSV 쓰기
# ==================================================

def write_csv_rows(
        path,
        rows,
        fieldnames
):

    with open(
            path,
            "w",
            encoding="utf-8-sig",
            newline=""
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


# ==================================================
# Excel 안전 문자열
# ==================================================

def safe_excel_value(value):

    if value is None:
        return ""

    s = str(value)

    # Excel에서 허용하지 않는
    # 제어문자 제거
    s = ILLEGAL_CHARACTERS_RE.sub(
        "",
        s
    )

    # Excel 셀 최대 글자수
    if len(s) > 32767:
        s = s[:32767]

    return s


# ==================================================
# CSV -> XLSX
# ==================================================

def csv_to_xlsx(
        csv_path,
        xlsx_path
):

    wb = Workbook()

    ws = wb.active

    ws.title = "data"


    with open(
            csv_path,
            "r",
            encoding="utf-8-sig",
            newline=""
    ) as f:

        reader = csv.reader(f)

        row_no = 1

        for row in reader:

            col_no = 1

            for value in row:

                cell = ws.cell(
                    row=row_no,
                    column=col_no,
                    value=safe_excel_value(
                        value
                    )
                )

                # ==================================================
                # 중요
                #
                # 연락처
                # 사업자번호
                # 숙소 ID
                #
                # 앞자리 0 손실 방지를 위해
                # 모두 문자열 형식
                # ==================================================

                cell.number_format = "@"

                col_no += 1

            row_no += 1


    # 첫 행 고정
    ws.freeze_panes = "A2"


    # 필터
    if (
            ws.max_row >= 1
            and ws.max_column >= 1
    ):

        ws.auto_filter.ref = (
            ws.dimensions
        )


    # ==================================================
    # 컬럼 너비 자동 조정
    # ==================================================

    max_width = 60

    for col_idx in range(
            1,
            ws.max_column + 1
    ):

        column_letter = (
            get_column_letter(
                col_idx
            )
        )

        width = 10

        # 너무 많은 행 전체 검사하면
        # 느리므로 앞 200행 기준
        for row_idx in range(
                1,
                min(
                    ws.max_row,
                    200
                ) + 1
        ):

            cell_value = (
                ws.cell(
                    row=row_idx,
                    column=col_idx
                ).value
            )

            if cell_value is not None:

                width = max(
                    width,
                    len(
                        str(
                            cell_value
                        )
                    ) + 2
                )

        if width > max_width:
            width = max_width

        ws.column_dimensions[
            column_letter
        ].width = width


    wb.save(
        xlsx_path
    )


# ==================================================
# Header
# ==================================================

def build_headers(stay_id):

    detail_url = (
            "https://nol.yanolja.com"
            "/stay/domestic/"
            + str(stay_id)
            + "?verticalCategory="
              "PRODUCT_CATEGORY_KOREA_ACCOMMODATION"
    )

    return {
        "accept": "*/*",

        "accept-language": (
            "ko-KR,ko;q=0.9,"
            "en-US;q=0.8,en;q=0.7"
        ),

        "platform": "Web",

        "referer": detail_url,

        "user-agent": (
            "Mozilla/5.0 "
            "(Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 "
            "(KHTML, like Gecko) "
            "Chrome/151.0.0.0 "
            "Safari/537.36"
        ),
    }


# ==================================================
# tRPC Params
# ==================================================
#
# TO-BE
#
# 0 -> getFavorite
# 1 -> getSellerInfo
#
# searchHome.home 제거됨
#
# verticalCategory payload도 제거됨
#
# ==================================================

def build_trpc_params(stay_id):

    payload = {
        "0": {
            "json": {
                "stayId": int(
                    stay_id
                )
            }
        },

        "1": {
            "json": {
                "stayId": int(
                    stay_id
                )
            }
        },
    }


    return {
        "batch": 1,

        "input": json.dumps(
            payload,
            ensure_ascii=False,
            separators=(",", ":")
        ),
    }


# ==================================================
# title 정규화
# ==================================================

def normalize_title(title):

    if title is None:
        return ""

    return (
        str(title)
        .replace(
            "\n",
            " "
        )
        .strip()
    )


# ==================================================
# 판매자정보 파싱
# ==================================================
#
# 응답
#
# [
#   0 = favorite
#   1 = sellerInfo
# ]
#
# ==================================================

def parse_seller_table(
        resp_json
):

    out = {}


    if (
            not resp_json
            or not isinstance(
        resp_json,
        list
    )
    ):

        return out


    # TO-BE는 최소 2개
    if len(resp_json) < 2:
        return out


    # ==================================================
    # index 1 = sellerInfo
    # ==================================================

    seller_response = (
            resp_json[1]
            or {}
    )


    result = (
            seller_response.get(
                "result"
            )
            or {}
    )


    data = (
            result.get(
                "data"
            )
            or {}
    )


    json_data = (
        data.get(
            "json"
        )
    )


    if (
            not json_data
            or not isinstance(
        json_data,
        list
    )
    ):

        return out


    # ==================================================
    # table block 찾기
    # ==================================================

    for block in json_data:

        if not isinstance(
                block,
                dict
        ):
            continue


        if (
                block.get("type")
                != "table"
        ):
            continue


        components = (
                block.get(
                    "tableComponent"
                )
                or []
        )


        # ==================================================
        # tableComponent
        # ==================================================

        for component in components:

            if not isinstance(
                    component,
                    dict
            ):
                continue


            title = normalize_title(
                component.get(
                    "title"
                )
            )


            bodys = (
                    component.get(
                        "bodys"
                    )
                    or []
            )


            # body가 배열
            if isinstance(
                    bodys,
                    list
            ):

                value = " | ".join(
                    [
                        str(x)
                        for x in bodys
                        if x is not None
                    ]
                )

            else:

                value = str(
                    bodys
                )


            if title:

                out[
                    title
                ] = value


        break


    return out


# ==================================================
# 숙소 1개 상세 호출
# ==================================================

def fetch_seller_info(
        stay_id,
        opt
):

    headers = build_headers(
        stay_id
    )

    params = build_trpc_params(
        stay_id
    )


    retry_count = (
        opt.get(
            "retry",
            3
        )
    )


    retry_sleep = (
        opt.get(
            "retrySleepSec",
            1
        )
    )


    timeout = (
        opt.get(
            "timeout",
            20
        )
    )


    last_err = None


    # Thread별 요청 안에서 Session 사용
    with requests.Session() as session:


        # 최초 + retry 횟수
        for attempt in range(
                1,
                retry_count + 2
        ):

            try:

                response = session.get(
                    TRPC_URL,
                    headers=headers,
                    params=params,
                    timeout=timeout
                )


                # ==================================================
                # 오류 로그
                # ==================================================

                if (
                        response.status_code
                        != 200
                ):

                    safe_print(
                        "[HTTP ERROR]"
                        + " | stayId={}".format(
                            stay_id
                        )
                        + " | status={}".format(
                            response.status_code
                        )
                        + " | attempt={}/{}".format(
                            attempt,
                            retry_count + 1
                        )
                        + " | body={}".format(
                            response.text[
                            :500
                            ]
                        )
                    )


                    response.raise_for_status()


                # ==================================================
                # JSON
                # ==================================================

                resp_json = (
                    response.json()
                )


                table_map = (
                    parse_seller_table(
                        resp_json
                    )
                )


                # API 200인데 판매자정보 없을 수도 있음
                if not table_map:

                    safe_print(
                        "[EMPTY SELLER]"
                        + " | stayId={}".format(
                            stay_id
                        )
                        + " | response={}".format(
                            str(
                                resp_json
                            )[:500]
                        )
                    )


                return {
                    "ok": True,

                    "stayId":
                        str(
                            stay_id
                        ),

                    "table":
                        table_map,

                    "err": ""
                }


            except Exception as e:

                last_err = str(
                    e
                )


                safe_print(
                    "[RETRY]"
                    + " | stayId={}".format(
                        stay_id
                    )
                    + " | attempt={}/{}".format(
                        attempt,
                        retry_count + 1
                    )
                    + " | error={}".format(
                        last_err
                    )
                )


                # 마지막이면 종료
                if attempt >= (
                        retry_count + 1
                ):

                    break


                wait_sec = (
                        retry_sleep
                        * attempt
                )


                safe_print(
                    "[RETRY WAIT]"
                    + " | stayId={}".format(
                        stay_id
                    )
                    + " | sleep={}sec".format(
                        wait_sec
                    )
                )


                time.sleep(
                    wait_sec
                )


    return {
        "ok": False,

        "stayId":
            str(
                stay_id
            ),

        "table": {},

        "err":
            last_err
            or "unknown error"
    }


# ==================================================
# 로그 값 줄이기
# ==================================================

def shorten_value(
        value,
        limit
):

    if value is None:
        return ""

    s = str(
        value
    )


    if (
            limit
            and limit > 0
            and len(s) > limit
    ):

        return (
                s[:limit]
                + "...(+"
                + str(
            len(s)
            - limit
        )
                + ")"
        )


    return s


# ==================================================
# 숙소 완료 로그
# ==================================================

def log_done(
        done_cnt,
        total_cnt,
        ok_cnt,
        stay_id,
        res,
        opt
):

    table = (
            res.get(
                "table"
            )
            or {}
    )


    keys = list(
        table.keys()
    )


    safe_print(
        "[DONE] {}/{}"
        " | ok={}/{}"
        " | stayId={}"
        " | fields={}{}"
        .format(
            done_cnt,
            total_cnt,

            ok_cnt,
            done_cnt,

            stay_id,

            len(keys),

            (
                ""
                if res.get("ok")
                else (
                        " | err="
                        + str(
                    res.get(
                        "err"
                    )
                )
                )
            )
        )
    )


    # ==================================================
    # 판매자정보 값 로그
    # ==================================================

    if table:

        limit = (
                opt.get(
                    "logValueLimit",
                    0
                )
                or 0
        )


        for key in table:

            value = (
                table.get(
                    key,
                    ""
                )
            )


            safe_print(
                "    - {} : {}".format(
                    key,
                    shorten_value(
                        value,
                        limit
                    )
                )
            )


# ==================================================
# 전체 CSV + 판매자정보 병합
# ==================================================

def merge_rows_with_seller(
        rows,
        opt
):

    tasks = []


    # ==================================================
    # 처리 대상 생성
    # ==================================================

    for index, row in enumerate(
            rows
    ):

        stay_id = (
            row.get(
                "id",
                ""
            )
        )


        stay_id = (
            str(
                stay_id
            )
            .strip()
        )


        if stay_id:

            tasks.append({
                "index":
                    index,

                "stayId":
                    stay_id
            })


    safe_print(
        "[START]"
        + " | input_rows={}".format(
            len(rows)
        )
        + " | tasks={}".format(
            len(tasks)
        )
        + " | workers={}".format(
            opt["maxWorkers"]
        )
    )


    # index 기준 결과
    seller_map_by_index = {}


    # ==================================================
    # 멀티쓰레드
    # ==================================================

    with ThreadPoolExecutor(
            max_workers=opt[
                "maxWorkers"
            ]
    ) as executor:


        future_map = {}


        for task in tasks:

            future = (
                executor.submit(
                    fetch_seller_info,

                    task[
                        "stayId"
                    ],

                    opt
                )
            )


            future_map[
                future
            ] = task


        done_cnt = 0

        ok_cnt = 0


        # ==================================================
        # 완료 순서대로
        # ==================================================

        for future in (
                as_completed(
                    future_map
                )
        ):


            meta = (
                future_map[
                    future
                ]
            )


            stay_id = (
                meta[
                    "stayId"
                ]
            )


            index = (
                meta[
                    "index"
                ]
            )


            try:

                result = (
                    future.result()
                )


            except Exception as e:

                result = {
                    "ok": False,

                    "stayId":
                        stay_id,

                    "table": {},

                    "err":
                        str(e)
                }


            done_cnt += 1


            if result.get(
                    "ok"
            ):

                ok_cnt += 1


            seller_map_by_index[
                index
            ] = result


            # ==================================================
            # 로그
            # ==================================================

            log_done(
                done_cnt,
                len(tasks),
                ok_cnt,
                stay_id,
                result,
                opt
            )


            if (
                    opt["sleepSec"]
                    and opt[
                "sleepSec"
            ] > 0
            ):

                time.sleep(
                    opt[
                        "sleepSec"
                    ]
                )


    # ==================================================
    # 원본 행과 병합
    # ==================================================

    out_rows = []

    all_new_cols = set()


    for index, original_row in enumerate(
            rows
    ):

        base = dict(
            original_row
        )


        result = (
            seller_map_by_index.get(
                index
            )
        )


        if (
                result
                and result.get(
            "table"
        )
        ):

            table = (
                    result.get(
                        "table"
                    )
                    or {}
            )


            for key in table.keys():

                all_new_cols.add(
                    key
                )

                base[
                    key
                ] = table.get(
                    key,
                    ""
                )


        out_rows.append(
            base
        )


    return (
        out_rows,
        sorted(
            list(
                all_new_cols
            )
        )
    )


# ==================================================
# 최종 컬럼 생성
# ==================================================

def build_fieldnames(
        original_rows,
        new_cols
):

    base_fields = []


    if original_rows:

        base_fields = list(
            original_rows[
                0
            ].keys()
        )


    union = set(
        base_fields
    )


    # ==================================================
    # 원본 CSV 중간 행에
    # 추가 컬럼이 있는 경우 대비
    # ==================================================

    for row in original_rows:

        for key in row.keys():

            if key not in union:

                union.add(
                    key
                )

                base_fields.append(
                    key
                )


    # ==================================================
    # 판매자 신규 컬럼
    # ==================================================

    for key in new_cols:

        if key not in base_fields:

            base_fields.append(
                key
            )


    return base_fields


# ==================================================
# 누락 컬럼 빈값
# ==================================================

def fill_missing_columns(
        rows,
        fieldnames
):

    for row in rows:

        for key in fieldnames:

            if key not in row:

                row[
                    key
                ] = ""


    return rows


# ==================================================
# MAIN
# ==================================================

def main():

    opt = (
        create_options()
    )


    safe_print(
        "=========================================="
    )

    safe_print(
        "야놀자 상세 판매자정보 수집 시작"
    )

    safe_print(
        "input={}".format(
            INPUT_CSV
        )
    )

    safe_print(
        "workers={}"
        " | retry={}"
        " | timeout={}"
        .format(
            opt[
                "maxWorkers"
            ],
            opt[
                "retry"
            ],
            opt[
                "timeout"
            ]
        )
    )

    safe_print(
        "=========================================="
    )


    # ==================================================
    # CSV 읽기
    # ==================================================

    rows = read_csv_rows(
        INPUT_CSV
    )


    if not rows:

        safe_print(
            "[EXIT]"
            + " | input csv empty"
            + " | path={}".format(
                INPUT_CSV
            )
        )

        return


    # ==================================================
    # 상세 수집
    # ==================================================

    out_rows, new_cols = (
        merge_rows_with_seller(
            rows,
            opt
        )
    )


    # ==================================================
    # 컬럼
    # ==================================================

    fieldnames = (
        build_fieldnames(
            rows,
            new_cols
        )
    )


    out_rows = (
        fill_missing_columns(
            out_rows,
            fieldnames
        )
    )


    # ==================================================
    # CSV 저장
    # ==================================================

    write_csv_rows(
        OUTPUT_CSV,
        out_rows,
        fieldnames
    )


    # ==================================================
    # XLSX 저장
    # ==================================================

    csv_to_xlsx(
        OUTPUT_CSV,
        OUTPUT_XLSX
    )


    safe_print(
        "=========================================="
    )


    safe_print(
        "[OK]"
        + " | output_rows={}".format(
            len(out_rows)
        )
        + " | csv={}".format(
            OUTPUT_CSV
        )
    )


    safe_print(
        "[OK]"
        + " | output_rows={}".format(
            len(out_rows)
        )
        + " | xlsx={}".format(
            OUTPUT_XLSX
        )
    )


    safe_print(
        "[OK]"
        + " | added_cols={}".format(
            len(new_cols)
        )
        + " | {}".format(
            ", ".join(
                new_cols
            )
        )
    )


    safe_print(
        "=========================================="
    )


if __name__ == "__main__":
    main()